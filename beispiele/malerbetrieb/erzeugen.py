#!/usr/bin/env python3
"""Erzeugt den synthetischen Belastungskorpus "Malerbetrieb Hoffmann 2025" in ./Quelle/.

Genau 200 Dateien in vier Unterordnern (Alt, Downloads, Fotos Handy, Scans),
deterministisch: feste Texte (texte.py = Dateien 1-100, eingefroren; texte2.py =
Dateien 101-200), reportlab `invariant=1`, feste Zip-Zeitstempel in den DOCX.
Rollen: ./README.md, Sollwerte: ./golden.json.
Aufruf: uv run python beispiele/malerbetrieb/erzeugen.py
"""
from __future__ import annotations

import io
import shutil
import sys
import textwrap
import zipfile
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

sys.path.insert(0, str(Path(__file__).resolve().parent))
import texte as T
import texte2 as T2

QUELLE = Path(__file__).resolve().parent / "Quelle"
LIEFERANTEN = {**T.LIEFERANTEN, **T2.LIEFERANTEN}
EINGANGSRECHNUNGEN = T.EINGANGSRECHNUNGEN + T2.EINGANGSRECHNUNGEN
AUSGANGSRECHNUNGEN = T.AUSGANGSRECHNUNGEN + T2.AUSGANGSRECHNUNGEN
ANGEBOTE = T.ANGEBOTE + T2.ANGEBOTE
ORDNER = ("Alt", "Downloads", "Fotos Handy", "Scans")
SCHRIFTDATEI = "/System/Library/Fonts/Supplemental/Arial.ttf"
SCAN_BREITE, SCAN_HOEHE, SCAN_PX = 1600, 2263, 40
# Schwelle ist scharf: ab ~5° Drehung findet Tesseract gar kein Wort mehr (Sicherheit 0.0),
# darunter sinkt die Wortkonfidenz mit der Unschärfe. Gemessene Werte: README.md.
SCAN_SCHLECHT_DREHUNG, SCAN_SCHLECHT_UNSCHAERFE = 3, 4.4
FOTO_BREITE, FOTO_PX = 1300, 72
FOTO_SAUBER_DREHUNG, FOTO_SCHLECHT_DREHUNG, FOTO_SCHLECHT_UNSCHAERFE = 1.5, 2.5, 7.3
FOTO_SAUBER_QUALITAET, FOTO_SCHLECHT_QUALITAET = 92, 85
ZIP_ZEIT = (2025, 1, 1, 0, 0, 0)
UST = Decimal("0.19")
CENT = Decimal("0.01")

# Ziel: Quelle -- bytegleiche Kopien, anderer Name, anderer Unterordner
DUPLIKATE: dict[str, str] = {
    "Scans/Rechnung_Farben_Maerz.pdf": "Downloads/Rechnung (4).pdf",
    "Downloads/Vertrag.pdf": "Alt/Rahmenvertrag_Geruestbau.pdf",
    "Scans/Beleg Baumarkt.jpg": "Fotos Handy/IMG_2231.jpg",
    "Scans/RE 2025-104 Kessler.pdf": "Downloads/2025-104.pdf",
    "Scans/Scan_0112.pdf": "Alt/Scan_0099.pdf",
}
RAHMENVERTRAG_PFAD = "Alt/Rahmenvertrag_Geruestbau.pdf"
URLAUBSFOTO_PFAD = "Fotos Handy/IMG_2650.jpg"
REZEPT_PFAD = "Downloads/rezept.txt"


# ----------------------------------------------------------------- Beträge

def eur(betrag: Decimal) -> str:
    """1234.5 -> '1.234,50'."""
    ganz, _, rest = f"{betrag.quantize(CENT):.2f}".partition(".")
    ganz_de = f"{int(ganz):,}".replace(",", ".")
    return f"{ganz_de},{rest}"


def summen(positionen: list[tuple]) -> tuple[Decimal, Decimal, Decimal]:
    """Netto, USt 19 %, Brutto -- kaufmännisch gerundet."""
    netto = sum((Decimal(str(m)) * Decimal(str(e)) for _, m, e in positionen), Decimal(0))
    netto = netto.quantize(CENT, ROUND_HALF_UP)
    ust = (netto * UST).quantize(CENT, ROUND_HALF_UP)
    return netto, ust, netto + ust


def positionszeilen(positionen: list[tuple]) -> list[str]:
    zeilen = [f"{'Pos Bezeichnung':42s}{'Menge':>5}   {'Einzel':>9}   {'Gesamt':>10}"]
    for nr, (bez, menge, einzel) in enumerate(positionen, start=1):
        gesamt = Decimal(str(menge)) * Decimal(str(einzel))
        zeilen.append(f"{nr:<3d} {bez[:38]:38s}{menge:>5}   {eur(Decimal(str(einzel))):>9}   {eur(gesamt):>10}")
    return zeilen


def summenzeilen(positionen: list[tuple], schluss: str) -> list[str]:
    netto, ust, brutto = summen(positionen)
    return [
        f"{'Nettobetrag':52s}{eur(netto):>12} EUR",
        f"{'Umsatzsteuer 19 %':52s}{eur(ust):>12} EUR",
        f"{schluss:52s}{eur(brutto):>12} EUR",
    ]


# ------------------------------------------------------------- Dokumenttexte

def rechnungsnummer(lieferant: str, n) -> str:
    return LIEFERANTEN[lieferant]["schema"].format(n=n)


def _rechnungsdatum(lieferant: str, n) -> str:
    for _, lf, nummer, datum, *_ in EINGANGSRECHNUNGEN:
        if lf == lieferant and nummer == n:
            return datum
    raise KeyError(f"Rechnung {lieferant} {n} fehlt")


def eingangsrechnung_text(r: tuple) -> str:
    _, lf, n, datum, betreff, positionen, korrektur_zu, grund = r
    lieferant = LIEFERANTEN[lf]
    nr = rechnungsnummer(lf, n)
    kopf = [f"{lieferant['name']}, {lieferant['adresse']}", ""]
    if korrektur_zu is None:
        kopf += [f"Rechnung Nr. {nr}", f"Rechnungsdatum {datum}, {lieferant['kunde']}"]
        schluss, fuss = "Rechnungsbetrag", []
    else:
        alt = rechnungsnummer(lf, korrektur_zu)
        kopf += [f"Korrekturrechnung Nr. {nr}", f"Rechnungsdatum {datum}, {lieferant['kunde']}",
                 f"zu Rechnung Nr. {alt} vom {_rechnungsdatum(lf, korrektur_zu)}", f"Grund der Korrektur: {grund}"]
        schluss, fuss = "Korrigierter Rechnungsbetrag", [f"Diese Korrekturrechnung ersetzt die Rechnung Nr. {alt}."]
    kopf += [f"Empfänger: {T.FIRMA}, {T.FIRMA_ADRESSE}", f"Betreff: {betreff}", ""]
    zahlung = ["", "Zahlbar innerhalb von 14 Tagen ohne Abzug.", lieferant["bank"]]
    return "\n".join(kopf + positionszeilen(positionen) + [""] + summenzeilen(positionen, schluss) + fuss + zahlung)


def _ausgangsrechnung(nr: str) -> tuple:
    for r in AUSGANGSRECHNUNGEN:
        if r[1] == nr:
            return r
    raise KeyError(f"Ausgangsrechnung {nr} fehlt")


def abzuege(teil) -> list[tuple[str, Decimal]]:
    """Schlussrechnung: (Zeilentext, Bruttobetrag) je bezahlter Teilrechnung -- mit Datum genannt,
    bewusst ohne Nummer, damit die Konfliktregel (Textverweis auf Rechnungsnummer) nicht anschlägt."""
    if not isinstance(teil, tuple):
        return []
    return [(f"abzüglich {i}. Teilrechnung vom {_ausgangsrechnung(nr)[2]}, bezahlt", summen(_ausgangsrechnung(nr)[7])[2])
            for i, nr in enumerate(teil[1], start=1)]


def restbetrag(r: tuple) -> Decimal | None:
    """Zu zahlender Rest einer Schlussrechnung, sonst None."""
    teil = r[8] if len(r) > 8 else None
    if not isinstance(teil, tuple):
        return None
    return summen(r[7])[2] - sum((b for _, b in abzuege(teil)), Decimal(0))


def ausgangsrechnung_text(r: tuple) -> str:
    _, nr, datum, kunde, kundenadresse, objekt, leistung, positionen, *rest = r
    teil = rest[0] if rest else None
    art = "Rechnung" if teil is None else (teil if isinstance(teil, str) else teil[0])
    kopf = [f"{T.FIRMA}, {T.FIRMA_ADRESSE}", "", f"{art} Nr. {nr}", f"Rechnungsdatum {datum}",
            f"Kunde: {kunde}, {kundenadresse}", f"Objekt: {objekt}", f"Leistung: {leistung}", ""]
    abzugszeilen = [f"{text:52s}{'-' + eur(brutto):>12} EUR" for text, brutto in abzuege(teil)]
    if abzugszeilen:
        abzugszeilen.append(f"{'Restbetrag zu zahlen':52s}{eur(restbetrag(r)):>12} EUR")
    fuss = ["", "Zahlbar innerhalb von 14 Tagen ohne Abzug. Vielen Dank für Ihren Auftrag.", "", T.FIRMA_FUSS]
    return "\n".join(kopf + positionszeilen(positionen) + [""] + summenzeilen(positionen, "Rechnungsbetrag") + abzugszeilen + fuss)


def paragraphen_seiten(kopf: str, paragraphen: list[tuple[str, str]], fueller: list[str]) -> list[str]:
    """Ein Paragraph je Seite, mit Fülltext auf Seitenlänge gebracht; Kopf nur auf Seite 1."""
    seiten = []
    for i, (titel, text) in enumerate(paragraphen):
        fuellsaetze = [fueller[(i + k) % len(fueller)] for k in range(4)]
        absaetze = [text] + [f"({k + 2}) {f}" for k, f in enumerate(fuellsaetze)]
        seitenkopf = kopf.strip("\n") + "\n\n" if i == 0 else ""
        seiten.append(seitenkopf + f"{titel}\n\n(1) " + "\n\n".join(absaetze) + f"\n\nSeite {i + 1} von {len(paragraphen)}")
    return seiten


def rahmenvertrag_seiten() -> list[str]:
    """30 Seiten Rahmenvertrag Gerüstbau (Datei 1-100, eingefroren)."""
    return paragraphen_seiten(T.RAHMENVERTRAG_KOPF, T.RAHMENVERTRAG_PARAGRAPHEN, T.RAHMENVERTRAG_FUELLER)


def versicherung_seiten() -> list[str]:
    """12 Seiten Betriebshaftpflicht-Police."""
    return paragraphen_seiten(T2.VERSICHERUNG_KOPF, T2.VERSICHERUNG_PARAGRAPHEN, T2.VERSICHERUNG_FUELLER)


# ------------------------------------------------------------------ Renderer

def schrift(px: int) -> ImageFont.FreeTypeFont:
    """Arial, weil die mitgelieferte Pillow-Schrift Umlaute und ß verfälscht."""
    try:
        return ImageFont.truetype(SCHRIFTDATEI, px)
    except OSError:
        return ImageFont.load_default(size=px)


def _umbrechen(text: str, breite: int = 88) -> list[str]:
    zeilen: list[str] = []
    for zeile in text.strip("\n").split("\n"):
        zeilen.extend(textwrap.wrap(zeile, breite) or [""])
    return zeilen


def text_pdf(pfad: Path, seiten: list[str]) -> None:
    """PDF mit echter Textebene (reportlab), eine Seite je Eintrag."""
    c = canvas.Canvas(str(pfad), pagesize=A4, invariant=1)
    for seite in seiten:
        c.setFont("Helvetica", 10)
        y = 800
        for zeile in _umbrechen(seite):
            c.drawString(50, y, zeile)
            y -= 15
        c.showPage()
    c.save()


def _malen(text: str, breite: int, hoehe: int, px: int) -> Image.Image:
    img = Image.new("RGB", (breite, hoehe), "white")
    d = ImageDraw.Draw(img)
    f = schrift(px)
    y = px
    for zeile in text.strip("\n").split("\n"):
        d.text((px, y), zeile, fill="black", font=f)
        y += int(px * 1.55)
    return img


def _verderben(bild: Image.Image, drehung: float, unschaerfe: float) -> Image.Image:
    bild = bild.rotate(drehung, fillcolor="white", expand=True)
    return bild.filter(ImageFilter.GaussianBlur(unschaerfe)) if unschaerfe else bild


def bild_pdf(pfad: Path, text: str, schlecht: bool = False) -> None:
    """Scan ohne Textebene: Bild als ganze A4-Seite. schlecht = gedreht + unscharf."""
    bild = _malen(text, SCAN_BREITE, SCAN_HOEHE, SCAN_PX)
    if schlecht:
        bild = _verderben(bild, SCAN_SCHLECHT_DREHUNG, SCAN_SCHLECHT_UNSCHAERFE)
    c = canvas.Canvas(str(pfad), pagesize=A4, invariant=1)
    c.drawImage(ImageReader(bild), 0, 0, width=A4[0], height=A4[1])
    c.save()


def handyfoto(pfad: Path, text: str, schlecht: bool) -> None:
    """Kassenbeleg als JPG: sauber = fast gerade, schlecht = schief + unscharf."""
    zeilen = text.strip("\n").count("\n") + 1
    hoehe = 2 * FOTO_PX + zeilen * int(FOTO_PX * 1.55)
    bild = _malen(text, FOTO_BREITE, hoehe, FOTO_PX)
    if schlecht:
        bild = _verderben(bild, FOTO_SCHLECHT_DREHUNG, FOTO_SCHLECHT_UNSCHAERFE)
        bild.save(pfad, "JPEG", quality=FOTO_SCHLECHT_QUALITAET)
    else:
        bild = _verderben(bild, FOTO_SAUBER_DREHUNG, 0)
        bild.save(pfad, "JPEG", quality=FOTO_SAUBER_QUALITAET)


def urlaubsfoto(pfad: Path) -> None:
    """Bild ohne Text: Himmel, Meer, Sonne."""
    breite, hoehe = 1600, 1200
    bild = Image.new("RGB", (breite, hoehe))
    d = ImageDraw.Draw(bild)
    for y in range(hoehe):
        if y < 700:
            farbe = (120 + y // 12, 170 + y // 20, 235)
        else:
            farbe = (30, 90 + (y - 700) // 8, 160)
        d.line([(0, y), (breite, y)], fill=farbe)
    d.ellipse([1150, 120, 1400, 370], fill=(255, 230, 120))
    bild.save(pfad, "JPEG", quality=88)


def angebot_docx(pfad: Path, a: tuple) -> None:
    from docx import Document

    _, nr, datum, kunde, objekt, positionen, hinweis = a
    doc = Document()
    doc.add_heading(f"Angebot Nr. {nr}", level=1)
    for zeile in (f"{T.FIRMA}, {T.FIRMA_ADRESSE}", f"Datum {datum}", f"Kunde: {kunde}", f"Objekt: {objekt}"):
        doc.add_paragraph(zeile)
    if hinweis:
        doc.add_paragraph(hinweis)
    tabelle = doc.add_table(rows=1, cols=4)
    for zelle, kopf in zip(tabelle.rows[0].cells, ("Bezeichnung", "Menge", "Einzelpreis", "Gesamt")):
        zelle.text = kopf
    for bez, menge, einzel in positionen:
        zeile = tabelle.add_row().cells
        gesamt = Decimal(str(menge)) * Decimal(str(einzel))
        for zelle, wert in zip(zeile, (bez, str(menge), eur(Decimal(str(einzel))), eur(gesamt))):
            zelle.text = wert
    netto, ust, brutto = summen(positionen)
    for zeile in (f"Nettobetrag {eur(netto)} EUR", f"Umsatzsteuer 19 % {eur(ust)} EUR",
                  f"Angebotssumme {eur(brutto)} EUR", "Dieses Angebot ist 30 Tage gültig. Ausführung nach Absprache."):
        doc.add_paragraph(zeile)
    doc.add_paragraph(T.FIRMA_FUSS)
    doc.save(str(pfad))
    _docx_normalisieren(pfad)


def _docx_normalisieren(pfad: Path) -> None:
    """Feste Zip-Zeitstempel, damit derselbe Inhalt bytegleich bleibt."""
    with zipfile.ZipFile(pfad) as alt:
        eintraege = [(info.filename, alt.read(info.filename)) for info in alt.infolist()]
    puffer = io.BytesIO()
    with zipfile.ZipFile(puffer, "w", zipfile.ZIP_DEFLATED) as neu:
        for name, daten in eintraege:
            neu.writestr(zipfile.ZipInfo(name, ZIP_ZEIT), daten, zipfile.ZIP_DEFLATED)
    pfad.write_bytes(puffer.getvalue())


# ------------------------------------------------------------------- Lauf

def _ziel(rel: str) -> Path:
    return QUELLE / rel


def erzeuge_rechnungen() -> None:
    for r in EINGANGSRECHNUNGEN:
        pfad, text = _ziel(r[0]), eingangsrechnung_text(r)
        if pfad.parent.name in ("Scans", "Alt") and pfad.name.startswith("Scan_"):
            bild_pdf(pfad, text)
        else:
            text_pdf(pfad, [text])
    for r in AUSGANGSRECHNUNGEN:
        text_pdf(_ziel(r[0]), [ausgangsrechnung_text(r)])
    for a in ANGEBOTE:
        angebot_docx(_ziel(a[0]), a)
    text_pdf(_ziel(T2.MAHNUNG[0]), [T2.MAHNUNG[1]])


def erzeuge_scans_und_fotos() -> None:
    for pfad, schlecht, text in T.KASSENBELEGE + T2.KASSENBELEGE:
        handyfoto(_ziel(pfad), text, schlecht)
    for pfad, schlecht, text in T.AUFMASSE + T2.AUFMASSE:
        bild_pdf(_ziel(pfad), text, schlecht)
    for pfad, text in T.BEHOERDEN_SCAN + T2.BEHOERDEN_SCAN:
        bild_pdf(_ziel(pfad), text)
    urlaubsfoto(_ziel(URLAUBSFOTO_PFAD))


def erzeuge_texte() -> None:
    text_pdf(_ziel(RAHMENVERTRAG_PFAD), rahmenvertrag_seiten())
    text_pdf(_ziel(T2.VERSICHERUNG_PFAD), versicherung_seiten())
    for pfad, seiten in T.VERTRAEGE + T.BEHOERDEN_TEXT + list(T.SACHFREMD_PDF.items()) + T2.VERTRAEGE + T2.BEHOERDEN_TEXT:
        text_pdf(_ziel(pfad), seiten)
    for pfad, text in {**T.NOTIZEN, **T.WIKI, REZEPT_PFAD: T.REZEPT, **T2.NOTIZEN, **T2.WIKI}.items():
        _ziel(pfad).write_text(text, encoding="utf-8")


def erzeuge_duplikate() -> None:
    for ziel, quelle in DUPLIKATE.items():
        shutil.copyfile(_ziel(quelle), _ziel(ziel))


def manifest() -> list[dict]:
    """Aussteller, Nummer, Datum, Bruttobetrag je Rechnung/Angebot -- für golden.json."""
    eintraege = []
    for pfad, lf, n, datum, _, positionen, korrektur_zu, _ in EINGANGSRECHNUNGEN:
        eintraege.append({"pfad": pfad, "art": "eingangsrechnung", "aussteller": LIEFERANTEN[lf]["name"],
                          "nummer": rechnungsnummer(lf, n), "datum": datum, "betrag": f"{eur(summen(positionen)[2])} EUR",
                          "korrektur_zu": None if korrektur_zu is None else rechnungsnummer(lf, korrektur_zu)})
    for r in AUSGANGSRECHNUNGEN:
        pfad, nr, datum, kunde, _, objekt, _, positionen, *rest = r
        teil = rest[0] if rest else None
        eintraege.append({"pfad": pfad, "art": "ausgangsrechnung", "aussteller": T.FIRMA, "nummer": nr, "datum": datum,
                          "betrag": f"{eur(summen(positionen)[2])} EUR", "kunde": kunde, "objekt": objekt,
                          "typ": "Rechnung" if teil is None else (teil if isinstance(teil, str) else teil[0]),
                          "restbetrag": None if restbetrag(r) is None else f"{eur(restbetrag(r))} EUR"})
    for pfad, nr, datum, kunde, objekt, positionen, _ in ANGEBOTE:
        eintraege.append({"pfad": pfad, "art": "angebot", "aussteller": T.FIRMA, "nummer": nr, "datum": datum,
                          "betrag": f"{eur(summen(positionen)[2])} EUR", "kunde": kunde, "objekt": objekt})
    return eintraege


def main() -> None:
    if QUELLE.exists():
        shutil.rmtree(QUELLE)
    for ordner in ORDNER:
        (QUELLE / ordner).mkdir(parents=True)
    erzeuge_rechnungen()
    erzeuge_scans_und_fotos()
    erzeuge_texte()
    erzeuge_duplikate()
    dateien = sorted(p for p in QUELLE.rglob("*") if p.is_file())
    for e in manifest():
        print(f"{e['pfad']:44s} {e['nummer']:16s} {e['datum']}  {e['betrag']:>14s}")
    print(f"\n{len(dateien)} Dateien, {sum(p.stat().st_size for p in dateien) / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
