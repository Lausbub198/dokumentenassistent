import sys
import types

import pytest
from docx import Document
from dokumentenassistent import auslesen
from dokumentenassistent.auslesen import (
    NichtLesbar,
    OcrEinstellungen,
    OcrNichtVerfuegbar,
    fordere_ocr,
    lies_seiten,
)

def test_M3_G1_text_pdf_liefert_eine_seite_je_pdf_seite(text_pdf):
    assert len(lies_seiten(text_pdf)) == 2

def test_M3_G2_text_pdf_enthaelt_text_ohne_ocr(text_pdf):
    s = lies_seiten(text_pdf)[0]
    assert ("2025-0815" in s.text, s.per_ocr, s.sicherheit) == (True, False, None)

def test_M3_G3_seitennummern_beginnen_bei_eins(text_pdf):
    assert [s.nummer for s in lies_seiten(text_pdf)] == [1, 2]

def test_M3_G4_bild_pdf_geht_durch_ocr(bild_pdf):
    s = lies_seiten(bild_pdf)[0]
    assert (s.per_ocr, "2025-0815" in s.text) == (True, True)

def test_M3_G5_sauberer_scan_hat_hohe_sicherheit(bild_pdf):
    assert lies_seiten(bild_pdf)[0].sicherheit >= 0.6

def test_M3_G6_jpg_liefert_eine_seite_per_ocr(unscharfes_bild):
    assert lies_seiten(unscharfes_bild)[0].per_ocr is True

def test_M3_B1_unscharfes_bild_hat_niedrige_sicherheit(unscharfes_bild):
    assert lies_seiten(unscharfes_bild)[0].sicherheit < 0.6

def test_M3_B2_unbekanntes_format_ergibt_nichtlesbar(tmp_path):
    p = tmp_path / "x.docx"; p.write_bytes(b"PK")
    with pytest.raises(NichtLesbar, match="docx"):
        lies_seiten(p)

def test_M3_B3_kaputtes_pdf_ergibt_nichtlesbar(tmp_path):
    p = tmp_path / "kaputt.pdf"; p.write_bytes(b"%PDF-1.4 garbage")
    with pytest.raises(NichtLesbar):
        lies_seiten(p)

def test_M3_B4_ohne_tesseract_ergibt_klaren_fehler(bild_pdf, monkeypatch):
    monkeypatch.setattr(auslesen, "tesseract_vorhanden", lambda: False)
    with pytest.raises(OcrNichtVerfuegbar, match="brew install tesseract"):
        lies_seiten(bild_pdf)

def test_M3_B5_leere_pdf_seite_ergibt_leeren_text(tmp_path):
    from tests.conftest import schreibe_text_pdf
    p = schreibe_text_pdf(tmp_path / "leer.pdf", [""])
    assert lies_seiten(p)[0].text.strip() == ""

def test_M3_G7_txt_liefert_eine_seite_ohne_ocr(tmp_path):
    p = tmp_path / "notiz.txt"; p.write_text("Hallo Welt", encoding="utf-8")
    seiten = lies_seiten(p)
    assert len(seiten) == 1
    s = seiten[0]
    assert (s.per_ocr, s.sicherheit, "Hallo Welt" in s.text) == (False, None, True)

def test_M3_G8_md_behaelt_markdown_syntax(tmp_path):
    p = tmp_path / "notiz.md"; p.write_text("# Titel\n\n**fett** Text", encoding="utf-8")
    s = lies_seiten(p)[0]
    assert "# Titel" in s.text and "**fett**" in s.text

def test_M3_G9_docx_absaetze_durch_leerzeilen_getrennt(tmp_path):
    d = Document()
    d.add_paragraph("Erster Absatz.")
    d.add_paragraph("Zweiter Absatz.")
    d.add_paragraph("Dritter Absatz.")
    p = tmp_path / "brief.docx"; d.save(p)
    text = lies_seiten(p)[0].text
    assert text.split("\n\n") == ["Erster Absatz.", "Zweiter Absatz.", "Dritter Absatz."]

def test_M3_G10_docx_tabelle_zeile_durch_pipe_verbunden(tmp_path):
    d = Document()
    tabelle = d.add_table(rows=1, cols=3)
    zellen = tabelle.rows[0].cells
    zellen[0].text, zellen[1].text, zellen[2].text = "Datum", "Betrag", "Aussteller"
    p = tmp_path / "tabelle.docx"; d.save(p)
    text = lies_seiten(p)[0].text
    assert "Datum | Betrag | Aussteller" in text

def test_M3_G11_lange_txt_wird_an_absaetzen_geteilt(tmp_path):
    absaetze = [f"Absatz {i} " + ("Wort " * 60) for i in range(20)]
    langer_text = "\n\n".join(absaetze)
    assert len(langer_text) > 3000
    p = tmp_path / "lang.txt"; p.write_text(langer_text, encoding="utf-8")
    seiten = lies_seiten(p)
    assert len(seiten) > 1
    assert [s.nummer for s in seiten] == list(range(1, len(seiten) + 1))
    assert all(len(s.text) <= 3000 for s in seiten)
    assert "".join(s.text for s in seiten).replace("\n\n", "") == langer_text.replace("\n\n", "")

def test_M3_B6_txt_latin1_fallback_liest_umlaute(tmp_path):
    p = tmp_path / "umlaut.txt"
    p.write_bytes("Prüfung bestanden, Größe ä".encode("latin-1"))
    s = lies_seiten(p)[0]
    assert "ä" in s.text

def test_M3_B7_leere_txt_ergibt_eine_seite_mit_leerem_text(tmp_path):
    p = tmp_path / "leer.txt"; p.write_text("", encoding="utf-8")
    seiten = lies_seiten(p)
    assert len(seiten) == 1 and seiten[0].text == ""

def test_M3_B8_kaputtes_docx_ergibt_nichtlesbar(tmp_path):
    p = tmp_path / "kaputt.docx"; p.write_bytes(b"PK garbage")
    with pytest.raises(NichtLesbar):
        lies_seiten(p)

def test_M3_B9_altes_doc_ergibt_nichtlesbar_mit_doc_in_meldung(tmp_path):
    p = tmp_path / "alt.doc"; p.write_bytes(b"\xd0\xcf\x11\xe0 garbage")
    with pytest.raises(NichtLesbar, match="doc"):
        lies_seiten(p)


# ---------------------------------------------------------------- OCR-Engine wählbar (24.09.2026)
#
# PaddleOCR wird in den Tests NIE geladen: das Paket zieht ein vollständiges Inferenz-Framework
# nach und lädt Modelle aus dem Netz. Ersetzt wird deshalb `auslesen.paddle_leser` — dieselbe
# Stelle, an der das Produkt den echten Leser holt. `test_M3_G17_...` wacht darüber.


class FakePaddle:
    """Ersatz für `PaddleOCR`: `predict` liefert die Struktur, die auch das echte Paket liefert."""

    def __init__(self, zeilen: list[str], scores: list[float]):
        self.zeilen = zeilen
        self.scores = scores
        self.bilder: list[str] = []

    def predict(self, bild: str):
        self.bilder.append(str(bild))
        return [{"rec_texts": list(self.zeilen), "rec_scores": list(self.scores)}]


def _paddle(monkeypatch, zeilen, scores) -> tuple[FakePaddle, list[tuple[str, bool]]]:
    """Setzt einen Fake-Leser ein und merkt sich, mit welchen (Sprache, Drehung) er geholt wurde."""
    leser = FakePaddle(zeilen, scores)
    gerufen: list[tuple[str, bool]] = []

    def fake_leser(sprache: str, drehung: bool):
        gerufen.append((sprache, drehung))
        return leser

    monkeypatch.setattr(auslesen, "paddle_leser", fake_leser)
    return leser, gerufen


PADDLE = OcrEinstellungen(engine="paddleocr", sprache="german", schwelle=0.80)


def test_M3_G12_standard_einstellungen_sind_das_heutige_verhalten():
    """Ohne Argument bleibt alles, wie es war — bestehende Aufrufer sind gültig."""
    e = OcrEinstellungen()
    assert (e.engine, e.sprache, e.schwelle) == ("tesseract", "deu", 0.6)


def test_M3_G13_paddleocr_liest_bild_und_mittelt_die_zeilenscores(unscharfes_bild, monkeypatch):
    _paddle(monkeypatch, ["Rechnung 2025-0815", "Betrag 1.240,00"], [0.99, 0.97])
    seite = lies_seiten(unscharfes_bild, PADDLE)[0]
    assert (seite.per_ocr, "2025-0815" in seite.text) == (True, True)
    assert seite.sicherheit == pytest.approx(0.98)


def test_M3_B10_paddleocr_ohne_erkennung_hat_sicherheit_null(unscharfes_bild, monkeypatch):
    """Die Regel „leeres Ergebnis ist kein sicheres Ergebnis" gilt für jede Engine."""
    _paddle(monkeypatch, [], [])
    seite = lies_seiten(unscharfes_bild, PADDLE)[0]
    assert (seite.text, seite.sicherheit) == ("", 0.0)


def test_M3_B11_tesseract_ohne_wortkasten_hat_sicherheit_null(unscharfes_bild, monkeypatch):
    monkeypatch.setattr(auslesen.pytesseract, "image_to_string", lambda *a, **k: "")
    monkeypatch.setattr(
        auslesen.pytesseract, "image_to_data", lambda *a, **k: {"conf": ["-1", "-1"]}
    )
    assert lies_seiten(unscharfes_bild)[0].sicherheit == 0.0


def test_M3_G14_drehung_und_sprache_gehen_an_den_paddle_leser(unscharfes_bild, monkeypatch):
    _, gerufen = _paddle(monkeypatch, ["x"], [0.9])
    lies_seiten(unscharfes_bild, OcrEinstellungen(engine="paddleocr", sprache="german", drehung=True))
    lies_seiten(unscharfes_bild, OcrEinstellungen(engine="paddleocr", sprache="german", drehung=False))
    assert gerufen == [("german", True), ("german", False)]


def test_M3_G15_tesseract_sprache_wird_durchgereicht(unscharfes_bild, monkeypatch):
    sprachen: list[str] = []
    monkeypatch.setattr(
        auslesen.pytesseract, "image_to_string",
        lambda bild, lang: sprachen.append(lang) or "",
    )
    monkeypatch.setattr(auslesen.pytesseract, "image_to_data", lambda *a, **k: {"conf": ["-1"]})
    lies_seiten(unscharfes_bild, OcrEinstellungen(sprache="eng"))
    assert sprachen == ["eng"]


def test_M3_G16_paddleocr_liest_auch_bild_pdfs(bild_pdf, monkeypatch):
    leser, _ = _paddle(monkeypatch, ["2025-0815"], [0.95])
    seite = lies_seiten(bild_pdf, PADDLE)[0]
    assert (seite.per_ocr, seite.text) == (True, "2025-0815")
    assert len(leser.bilder) == 1


def test_M3_B12_paddleocr_gewaehlt_aber_paket_fehlt_ergibt_klaren_fehler(monkeypatch):
    """Kein stiller Rückfall auf Tesseract: fehlt das Extra, bricht es mit Hinweis ab."""
    monkeypatch.setattr(auslesen, "paddleocr_vorhanden", lambda: False)
    with pytest.raises(OcrNichtVerfuegbar) as exc:
        fordere_ocr(PADDLE)
    assert "uv sync --extra ocr-paddle" in str(exc.value)


def test_M3_B13_unbekannte_engine_ergibt_klaren_fehler(unscharfes_bild):
    with pytest.raises(OcrNichtVerfuegbar, match="surya"):
        lies_seiten(unscharfes_bild, OcrEinstellungen(engine="surya"))


def test_M3_G17_paddleocr_wird_in_den_tests_nie_geladen():
    """Wächter: schlägt an, sobald ein Test das echte Paket importiert (Modelle aus dem Netz)."""
    assert "paddleocr" not in sys.modules


def _paddle_kwargs(monkeypatch, drehung: bool) -> dict:
    """Ruft `paddle_leser` gegen ein Fake-Modul `paddleocr` und gibt die Konstruktor-Argumente.

    `monkeypatch.setitem` nimmt den Eintrag in `sys.modules` nach dem Test wieder heraus —
    der Wächter `test_M3_G17_...` bleibt gültig, das echte Paket wird nie geladen.
    """
    gesehen: dict = {}

    class FakePaddleOCR:
        def __init__(self, **kwargs):
            gesehen.update(kwargs)

    modul = types.ModuleType("paddleocr")
    modul.PaddleOCR = FakePaddleOCR
    monkeypatch.setitem(sys.modules, "paddleocr", modul)
    monkeypatch.setattr(auslesen, "_PADDLE_LESER", {})
    auslesen.paddle_leser("german", drehung)
    return gesehen


def test_M3_G18_drehung_schaltet_nur_die_seitenorientierung(monkeypatch):
    """Gemessen 24.09.2026 (beweise/ocr-drehung/2026-09-24/): nur diese Stufe hilft.

    Der Sammelschalter „alle drei an" liest über die zehn Seiten 168 Zahlen und damit
    WENIGER als gar keine Drehung (179); allein die Seitenorientierung liest 182.
    Entzerrung und Zeilenorientierung bleiben deshalb aus — auch wenn `drehung` an ist.
    """
    assert _paddle_kwargs(monkeypatch, True) == {
        "lang": "german",
        "use_doc_orientation_classify": True,
        "use_doc_unwarping": False,
        "use_textline_orientation": False,
    }


def test_M3_G19_ohne_drehung_ist_keine_vorstufe_an(monkeypatch):
    assert _paddle_kwargs(monkeypatch, False) == {
        "lang": "german",
        "use_doc_orientation_classify": False,
        "use_doc_unwarping": False,
        "use_textline_orientation": False,
    }
